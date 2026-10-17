"""Build briefings/housing-market-briefing.pptx from the CSV files in data/.

Usage:  .venv/bin/python briefings/build_deck.py [--render]
--render converts the deck to PDF with LibreOffice and writes one PNG per slide to briefings/png/.
"""
import csv, shutil, subprocess, sys
from collections import defaultdict
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.chart import XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.parts.slide import NotesMasterPart
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
TEMPLATE = ROOT / "templates" / "github-dark-template.pptx"
OUT = ROOT / "briefings" / "housing-market-briefing.pptx"
MONTHS = "Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split()
DISCLAIMER = "Built from public FHFA and CFPB data. Not an official Fannie Mae, FHFA, or CFPB product."
FONT = "Arial"

# Colors from the GitHub dark template. Blue and gold are a color-blind-safe pair.
BG, CARD, BORDER = RGBColor(0x0D, 0x11, 0x17), RGBColor(0x16, 0x1B, 0x22), RGBColor(0x30, 0x36, 0x3D)
WHITE, BODY, MUTED = RGBColor(0xF0, 0xF6, 0xFC), RGBColor(0xB1, 0xBA, 0xC4), RGBColor(0x8B, 0x94, 0x9E)
BLUE, GOLD, GREEN = RGBColor(0x58, 0xA6, 0xFF), RGBColor(0xD2, 0x99, 0x22), RGBColor(0x3F, 0xB9, 0x50)


# ---------- data ----------
def rows(name):
    with open(DATA / name, newline="") as f:
        return list(csv.DictReader(f))


def num(v):
    return float(v) if v not in ("", None) else None


def pct(x, signed=True):
    return f"{x * 100:+.1f}%" if signed else f"{abs(x) * 100:.1f}%"


def mon(ym):
    return f"{MONTHS[ym[1] - 1]} {ym[0]}"


def qtr(yq):
    return f"Q{yq[1]} {yq[0]}"


def short_metro(name):
    city, _, states = name.partition(",")
    return f"{city.split('-')[0].strip()}, {states.strip()[:2]}"


def quarterly_yoy(file, key, label):
    idx, names = defaultdict(dict), {}
    for r in rows(file):
        v = num(r["index_nsa"])
        if v:
            idx[r[key]][(int(r["year"]), int(r["quarter"]))] = v
            names[r[key]] = label(r)
    lq = max(k for d in idx.values() for k in d)
    prev = (lq[0] - 1, lq[1])
    yoy = {names[c]: d[lq] / d[prev] - 1 for c, d in idx.items() if lq in d and prev in d}
    return lq, sorted(yoy.items(), key=lambda kv: kv[1])


def load():
    d = {}
    us = [r for r in rows("hpi_us_division_monthly.csv") if r["place_id"] == "USA"]
    sa = {(int(r["year"]), int(r["month"])): num(r["index_sa"]) for r in us if num(r["index_sa"])}
    nsa = {(int(r["year"]), int(r["month"])): num(r["index_nsa"]) for r in us if num(r["index_nsa"])}
    yoy = {k: v / nsa[(k[0] - 1, k[1])] - 1 for k, v in nsa.items() if (k[0] - 1, k[1]) in nsa}
    latest, peak = max(sa), max(sa, key=sa.get)
    old_peak = max((k for k in sa if 2005 <= k[0] <= 2008), key=sa.get)
    trough = min((k for k in sa if 2009 <= k[0] <= 2013), key=sa.get)
    yoy_peak = max((k for k in yoy if k[0] >= 2020), key=yoy.get)
    pre = [v for k, v in yoy.items() if 2015 <= k[0] <= 2019]
    d.update(sa=sa, yoy=yoy, latest=latest, latest_sa=sa[latest], peak=peak, from_peak=sa[latest] / sa[peak] - 1,
             since2020=sa[latest] / sa[(2020, 1)] - 1, old_peak=old_peak, trough=trough,
             bust=sa[trough] / sa[old_peak] - 1, vs_old_peak=sa[latest] / sa[old_peak] - 1,
             latest_yoy=yoy[max(yoy)], latest_yoy_m=max(yoy), yoy_peak=yoy_peak, yoy_peak_v=yoy[yoy_peak],
             pre_avg=sum(pre) / len(pre))
    d["level"] = ("at a record high" if latest == peak else
                  "near their record high" if d["from_peak"] > -0.01 else
                  f"{pct(d['from_peak'], False)} below their {mon(peak)} peak")

    d["sq"], d["states"] = quarterly_yoy("hpi_state_quarterly.csv", "state_code", lambda r: r["state_name"])
    d["mq"], d["metros"] = quarterly_yoy("hpi_metro_quarterly.csv", "cbsa_code", lambda r: short_metro(r["metro_name"]))

    loans, vol = defaultdict(lambda: defaultdict(float)), defaultdict(lambda: defaultdict(float))
    for r in rows("hmda_state_originations.csv"):
        loans[int(r["year"])][r["loan_purpose"]] += float(r["loans"])
        vol[int(r["year"])][r["loan_purpose"]] += float(r["dollar_volume"])
    years = sorted(loans)
    purchase = [loans[y]["Home purchase"] for y in years]
    refi = [loans[y]["Refinance"] + loans[y]["Cash-out refinance"] for y in years]
    pk = max(range(len(years)), key=lambda i: refi[i])
    avg = lambda y: vol[y]["Home purchase"] / loans[y]["Home purchase"]
    d.update(years=years, purchase=purchase, refi=refi, refi_peak_i=pk,
             refi_from_peak=refi[-1] / refi[pk] - 1, purchase_vs_first=purchase[-1] / purchase[0] - 1,
             purchase_share=purchase[-1] / (purchase[-1] + refi[-1]),
             purchase_share_peak=purchase[pk] / (purchase[pk] + refi[pk]),
             avg_first=avg(years[0]), avg_last=avg(years[-1]))
    return d


# ---------- drawing helpers ----------
class Deck:
    def __init__(self):
        self.prs = Presentation(TEMPLATE)
        ids = self.prs.slides._sldIdLst
        for s in list(ids):  # keep the template's master, layout and size; drop its sample slides
            self.prs.part.drop_rel(s.rId)
            ids.remove(s)
        self.layout = self.prs.slide_layouts[0]
        # The template's notes master has no notes placeholder, so swap in python-pptx's default one.
        part = self.prs.part
        old = next(r for r in part.rels.values() if r.reltype == RT.NOTES_MASTER)
        part.drop_rel(old.rId)
        new_rid = part.relate_to(NotesMasterPart.create_default(part.package), RT.NOTES_MASTER)
        self.prs._element.find(qn("p:notesMasterIdLst"))[0].set(qn("r:id"), new_rid)

    def new(self, notes):
        sl = self.prs.slides.add_slide(self.layout)
        for ph in list(sl.placeholders):
            ph._element.getparent().remove(ph._element)
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = BG
        sl.notes_slide.notes_text_frame.text = notes
        return sl

    def text(self, sl, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
        """paras: list of (text, size, color, bold)."""
        tf = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
        tf.word_wrap, tf.vertical_anchor = True, anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for n, (t, size, color, bold) in enumerate(paras):
            p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
            if n:
                p.space_before = Pt(size * 0.4)
            r = p.add_run()
            r.text = t
            r.font.size, r.font.color.rgb, r.font.bold, r.font.name = Pt(size), color, bold, FONT
        return tf

    def rect(self, sl, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
        r = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        r.fill.solid()
        r.fill.fore_color.rgb = color
        r.line.fill.background()
        r.shadow.inherit = False
        return r

    def card(self, sl, x, y, w, h, accent=None):
        c = self.rect(sl, x, y, w, h, CARD, MSO_SHAPE.ROUNDED_RECTANGLE)
        c.adjustments[0] = 0.04
        c.line.color.rgb, c.line.width = BORDER, Pt(1)
        if accent:  # template cards carry a colored bar across the top edge
            self.rect(sl, x, y, w, 0.06, accent)

    def header(self, sl, eyebrow, headline):
        self.text(sl, 0.45, 0.3, 9.1, 0.25, [(eyebrow.upper(), 10.5, BLUE, True)])
        self.text(sl, 0.45, 0.56, 9.1, 0.8, [(headline, 22, WHITE, True)])

    def footer(self, sl, source):
        self.text(sl, 0.45, 5.1, 9.1, 0.4, [(source, 8, MUTED, False), (DISCLAIMER, 8, MUTED, False)])

    def chart(self, sl, kind, cd, x=0.45, y=1.45, w=9.1, h=3.5):
        self.card(sl, x, y, w, h)
        return sl.shapes.add_chart(kind, Inches(x + 0.1), Inches(y + 0.08), Inches(w - 0.2), Inches(h - 0.16), cd).chart


def _nofill():
    sp = etree.Element(qn("c:spPr"))
    etree.SubElement(sp, qn("a:noFill"))
    etree.SubElement(etree.SubElement(sp, qn("a:ln")), qn("a:noFill"))
    return sp


def _font(tf, size, color, bold):
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size, r.font.color.rgb, r.font.bold, r.font.name = Pt(size), color, bold, FONT


def style_chart(chart, title, x_title, y_title, fmt, legend=False):
    chart.font.size, chart.font.color.rgb, chart.font.name = Pt(10), BODY, FONT
    c = chart._chartSpace.find(qn("c:chart"))
    c.addnext(_nofill())  # transparent chart area so the card shows through
    c.find(qn("c:plotArea")).append(_nofill())
    chart.has_title = True
    chart.chart_title.text_frame.text = title
    _font(chart.chart_title.text_frame, 12, WHITE, True)
    ca, va = chart.category_axis, chart.value_axis
    for ax, t in ((ca, x_title), (va, y_title)):
        ax.has_title = True
        ax.axis_title.text_frame.text = t
        _font(ax.axis_title.text_frame, 10, MUTED, False)
        ax.tick_labels.font.size, ax.tick_labels.font.color.rgb = Pt(10), BODY
        ax.major_tick_mark = ax.minor_tick_mark = XL_TICK_MARK.NONE
    ca.format.line.color.rgb = BORDER
    va.format.line.fill.background()
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = BORDER
    va.major_gridlines.format.line.width = Pt(0.75)
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = fmt, False
    chart.has_legend = legend
    if legend:
        chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.TOP, False
        chart.legend.font.size, chart.legend.font.color.rgb = Pt(10), BODY
    return ca, va


def fill(series, color):
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = color


def skip_labels(ax, n):
    anchor = ax._element.find(qn("c:noMultiLvlLbl"))
    for tag in ("c:tickLblSkip", "c:tickMarkSkip"):
        e = etree.Element(qn(tag))
        e.set("val", str(n))
        anchor.addprevious(e)


def label_point(series, i, text, pos):
    dl = series.points[i].data_label
    dl.text_frame.text = text
    _font(dl.text_frame, 10, WHITE, True)
    dl.position = pos


def value_labels(plot, fmt):
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = fmt, False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.color.rgb = Pt(9), BODY


def stat_cards(dk, sl, y, h, items, big_size, title_size, body_size):
    w = (9.1 - 2 * 0.2) / 3
    for n, (big, title, body, color) in enumerate(items):
        x = 0.45 + n * (w + 0.2)
        dk.card(sl, x, y, w, h, accent=color)
        dk.text(sl, x + 0.22, y + 0.25, w - 0.44, big_size / 60, [(big, big_size, color, True)])
        paras = ([(title, title_size, WHITE, True)] if title else []) + [(body, body_size, BODY, False)]
        top = y + 0.35 + big_size / 60
        dk.text(sl, x + 0.22, top, w - 0.44, y + h - top - 0.15, paras)


# ---------- slides ----------
def build(d):
    dk = Deck()
    lm, sq, mq, years = d["latest"], d["sq"], d["mq"], d["years"]
    states, metros = d["states"], d["metros"]
    s_down, m_down = sum(v < 0 for _, v in states), sum(v < 0 for _, v in metros)
    ly, py, pk = years[-1], years[d["refi_peak_i"]], d["refi_peak_i"]
    hpi_src = f"Source: FHFA House Price Index, purchase-only. Latest period: {mon(lm)}."

    # 1. Title and key numbers
    sl = dk.new(
        f"This briefing uses public FHFA house price data through {mon(lm)} and CFPB HMDA mortgage data for "
        f"{years[0]} to {ly}. The short story: U.S. home prices are {d['level']} and {pct(d['since2020'])} above "
        f"January 2020. But annual price growth slowed to {pct(d['latest_yoy'])}, down from {pct(d['yoy_peak_v'])} "
        f"in {mon(d['yoy_peak'])}. Refinance loans fell {pct(d['refi_from_peak'], False)} from their {py} peak. "
        "The next slides show the price level, the speed of change, the states and metros that lead and lag, "
        "mortgage originations, and three takeaways. All numbers are facts from the data, not forecasts.")
    dk.rect(sl, 0.45, 0.55, 0.6, 0.06, GREEN)
    dk.text(sl, 0.45, 0.3, 9.1, 0.25, [("EXECUTIVE BRIEFING  |  U.S. HOUSING MARKET", 10.5, GREEN, True)])
    dk.text(sl, 0.45, 0.8, 9.1, 1.25, [(f"U.S. home prices are {d['level']}, but annual growth slowed "
                                        f"to {pct(d['latest_yoy'])}", 30, WHITE, True)])
    dk.text(sl, 0.45, 2.15, 9.1, 0.3, [(f"Home prices through {mon(lm)} (FHFA HPI)  |  Mortgage originations "
                                        f"{years[0]} to {ly} (CFPB HMDA)", 14, BODY, False)])
    stat_cards(dk, sl, 2.8, 1.95, [
        (pct(d["since2020"]), None, f"Change in U.S. home prices, Jan 2020 to {mon(lm)}", BLUE),
        (pct(d["latest_yoy"]), None, f"Annual price change in {mon(d['latest_yoy_m'])}, down from "
                                     f"{pct(d['yoy_peak_v'])} in {mon(d['yoy_peak'])}", GOLD),
        (pct(d["refi_from_peak"]), None, f"Change in refinance loans, {py} peak to {ly}", GREEN)], 36, 0, 14)
    dk.footer(sl, f"Sources: FHFA House Price Index, purchase-only (latest {mon(lm)}); "
                  f"CFPB HMDA originations (latest {ly}).")

    # 2. Price level
    sl = dk.new(
        "The FHFA index measures repeat sales of the same homes, so it tracks price change, not a dollar price. "
        f"In {mon(lm)} the seasonally adjusted U.S. index was {d['latest_sa']:.0f}, where 1991 Q1 equals 100. "
        f"That is {pct(d['since2020'])} above January 2020 and {pct(d['vs_old_peak'])} above the "
        f"{mon(d['old_peak'])} peak before the housing crash. For context, the index fell "
        f"{pct(d['bust'], False)} from that peak to its {mon(d['trough'])} low. Today, prices are {d['level']}.")
    dk.header(sl, "01  Where prices are now",
              f"U.S. home prices are {d['level']} and {pct(d['since2020'], False)} above January 2020")
    keys = sorted(k for k in d["sa"] if k[0] >= 2000)
    cd = CategoryChartData()
    cd.categories = [mon(k) for k in keys]
    cd.add_series("U.S. house price index (SA)", [round(d["sa"][k], 1) for k in keys])
    ch = dk.chart(sl, XL_CHART_TYPE.LINE, cd)
    ca, _ = style_chart(ch, "FHFA U.S. House Price Index, seasonally adjusted, monthly, 2000 to latest",
                        "Month", "Index (1991 Q1 = 100)", "0")
    skip_labels(ca, 36)
    ser = ch.plots[0].series[0]
    ser.smooth = False
    ser.marker.style = XL_MARKER_STYLE.NONE
    ser.format.line.color.rgb, ser.format.line.width = BLUE, Pt(2.25)
    label_point(ser, keys.index(d["old_peak"]), f"{mon(d['old_peak'])}: {d['sa'][d['old_peak']]:.0f}",
                XL_LABEL_POSITION.ABOVE)
    label_point(ser, keys.index(d["trough"]), f"{mon(d['trough'])}: {d['sa'][d['trough']]:.0f}",
                XL_LABEL_POSITION.BELOW)
    label_point(ser, len(keys) - 1, f"{mon(lm)}: {d['latest_sa']:.0f}", XL_LABEL_POSITION.ABOVE)
    dk.footer(sl, hpi_src + " Seasonally adjusted index, 1991 Q1 = 100.")

    # 3. Speed of change
    sl = dk.new(
        "This chart shows the change in prices from the same month one year earlier, using the non-seasonally "
        f"adjusted index. From 2015 to 2019, annual growth averaged about {pct(d['pre_avg'])}. Growth peaked at "
        f"{pct(d['yoy_peak_v'])} in {mon(d['yoy_peak'])}. In {mon(d['latest_yoy_m'])} it was "
        f"{pct(d['latest_yoy'])}. A positive but smaller number means prices still rose, only more slowly. "
        "The data does not tell us what happens next.")
    dk.header(sl, "02  How fast prices changed",
              f"Annual price growth slowed from {pct(d['yoy_peak_v'])} in {mon(d['yoy_peak'])} "
              f"to {pct(d['latest_yoy'])} in {mon(d['latest_yoy_m'])}")
    keys = sorted(k for k in d["yoy"] if k[0] >= 2016)
    cd = CategoryChartData()
    cd.categories = [mon(k) for k in keys]
    cd.add_series("Year-over-year change", [round(d["yoy"][k], 6) for k in keys])
    ch = dk.chart(sl, XL_CHART_TYPE.COLUMN_CLUSTERED, cd)
    ca, _ = style_chart(ch, "U.S. home prices: change from the same month one year earlier",
                        "Month", "Year-over-year change", "0%")
    skip_labels(ca, 12)
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    ch.plots[0].gap_width = 40
    ser = ch.plots[0].series[0]
    ser.invert_if_negative = False
    fill(ser, BLUE)
    va = ch.value_axis  # fixed 0 to peak + headroom, in 4% steps, so the peak label fits
    va.minimum_scale, va.major_unit = 0, 0.04
    va.maximum_scale = (int(d["yoy_peak_v"] / 0.04) + 2) * 0.04
    for k in (d["yoy_peak"], keys[-1]):
        i = keys.index(k)
        fill(ser.points[i], GOLD)
        label_point(ser, i, pct(d["yoy"][k]), XL_LABEL_POSITION.OUTSIDE_END)
    dk.footer(sl, hpi_src + " Year-over-year change uses the non-seasonally adjusted index.")

    # 4 and 5. Leaders and laggards
    def leaders(eyebrow, items, period, n_down, noun, src_note, headline):
        total, prev = len(items), (period[0] - 1, period[1])
        sl = dk.new(
            f"This chart ranks {total} {noun} by the change in prices from {qtr(prev)} to {qtr(period)}. "
            f"The five fastest were {', '.join(f'{k} ({pct(v)})' for k, v in reversed(items[-5:]))}. "
            f"The five slowest were {', '.join(f'{k} ({pct(v)})' for k, v in items[:5])}. "
            f"{n_down} of {total} {noun} had a price decline over the year. "
            "These are index changes for each area, not dollar prices.")
        dk.header(sl, eyebrow, headline)
        pick = items[:5] + items[-5:]
        cd = CategoryChartData()
        cd.categories = [k for k, _ in pick]
        cd.add_series("5 fastest", [None] * 5 + [round(v, 6) for _, v in pick[5:]])
        cd.add_series("5 slowest", [round(v, 6) for _, v in pick[:5]] + [None] * 5)
        ch = dk.chart(sl, XL_CHART_TYPE.BAR_CLUSTERED, cd)
        ca, _ = style_chart(ch, f"{noun.capitalize()}: 5 fastest and 5 slowest annual price change, {qtr(period)}",
                            noun.capitalize()[:-1], "Year-over-year change", "0%", legend=True)
        ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        plot = ch.plots[0]
        plot.overlap, plot.gap_width = 100, 40
        for ser, color in zip(plot.series, (BLUE, GOLD)):
            ser.invert_if_negative = False
            fill(ser, color)
        value_labels(plot, "+0.0%;-0.0%;0.0%")
        dk.footer(sl, f"Source: FHFA House Price Index, purchase-only, {src_note}. Latest period: {qtr(period)}. "
                      f"Change vs. {qtr(prev)}, non-seasonally adjusted.")

    (top_s, top_sv), (bot_s, bot_sv) = states[-1], states[0]
    leaders("03  States that lead and lag", states, sq, s_down, "states", "50 states and DC",
            f"{top_s} leads with {pct(top_sv)} a year; {bot_s} lags with {pct(bot_sv)}")
    (top_m, top_mv), (bot_m, bot_mv) = metros[-1], metros[0]
    leaders("04  Metros that lead and lag", metros, mq, m_down, "metros", "100 largest metros",
            f"{top_m} leads at {pct(top_mv)}; prices fell in {m_down} of {len(metros)} metros")

    # 6. Originations
    sl = dk.new(
        "These are counts of originated loans reported under HMDA, summed across all states. Refinance combines "
        f"rate-and-term refinance and cash-out refinance. Refinance loans peaked at {d['refi'][pk] / 1e6:.1f} "
        f"million in {py} and were {d['refi'][-1] / 1e6:.1f} million in {ly}, a change of {pct(d['refi_from_peak'])}. "
        f"Home purchase loans went from {d['purchase'][0] / 1e6:.1f} million in {years[0]} to "
        f"{d['purchase'][-1] / 1e6:.1f} million in {ly} ({pct(d['purchase_vs_first'])}). Purchase loans were "
        f"{pct(d['purchase_share_peak'], False)} of originations in {py} and {pct(d['purchase_share'], False)} in {ly}. "
        f"The average home purchase loan grew from ${d['avg_first'] / 1e3:,.0f}K to ${d['avg_last'] / 1e3:,.0f}K.")
    dk.header(sl, "05  Mortgage originations",
              f"Refinance loans fell {pct(d['refi_from_peak'], False)} from their {py} peak; home purchase "
              f"loans were {pct(d['purchase_share'], False)} of {ly} originations")
    cd = CategoryChartData()
    cd.categories = [str(y) for y in years]
    cd.add_series("Home purchase", [round(v / 1e6, 6) for v in d["purchase"]])
    cd.add_series("Refinance (incl. cash-out)", [round(v / 1e6, 6) for v in d["refi"]])
    ch = dk.chart(sl, XL_CHART_TYPE.COLUMN_CLUSTERED, cd)
    style_chart(ch, f"Originated mortgage loans by purpose, {years[0]} to {ly}", "Year", "Loans (millions)",
                "0.0", legend=True)
    plot = ch.plots[0]
    plot.gap_width, plot.overlap = 60, -10
    for ser, color in zip(plot.series, (BLUE, GOLD)):
        fill(ser, color)
    value_labels(plot, "0.0")
    dk.footer(sl, f"Source: CFPB HMDA Data Browser, originated loans (action taken = 1), all states. "
                  f"Latest period: {ly}.")

    # 7. Takeaways
    sl = dk.new(
        f"Three facts to remember. One: prices are {d['level']} and {pct(d['since2020'])} above January 2020. "
        f"Two: annual growth slowed to {pct(d['latest_yoy'])}, and in {qtr(sq)} {s_down} states and {m_down} large "
        f"metros had price declines. Three: refinance loans are {pct(d['refi_from_peak'], False)} below their {py} "
        f"peak, so home purchase loans were {pct(d['purchase_share'], False)} of {ly} originations. "
        "These are facts from the data, not forecasts.")
    dk.header(sl, "06  Three takeaways", "Prices are high, growth has slowed, and refinancing fell")
    stat_cards(dk, sl, 1.45, 3.5, [
        (pct(d["since2020"]), "Prices are high",
         f"The U.S. index is {pct(d['since2020'], False)} above January 2020 and {d['level']} in {mon(lm)}.", BLUE),
        (pct(d["latest_yoy"]), "Growth has slowed",
         f"Annual growth fell from {pct(d['yoy_peak_v'])} in {mon(d['yoy_peak'])}. In {qtr(sq)}, {s_down} states "
         f"and {m_down} of {len(metros)} large metros had price declines.", GOLD),
        (pct(d["refi_from_peak"]), "Refinancing fell",
         f"Refinance loans fell from {d['refi'][pk] / 1e6:.1f}M in {py} to {d['refi'][-1] / 1e6:.1f}M in {ly}. "
         f"Purchase loans were {pct(d['purchase_share'], False)} of {ly} originations.", GREEN)], 44, 18, 15)
    dk.footer(sl, f"Sources: FHFA House Price Index, purchase-only (latest {mon(lm)}; {qtr(sq)} for states and "
                  f"metros); CFPB HMDA originations (latest {ly}).")

    OUT.parent.mkdir(exist_ok=True)
    dk.prs.save(OUT)
    print(f"Saved {OUT.relative_to(ROOT)} ({len(dk.prs.slides)} slides)")


def render():
    import fitz  # PyMuPDF
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    out = OUT.parent / "png"
    out.mkdir(exist_ok=True)
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out), str(OUT)],
                   check=True, capture_output=True)
    pdf = out / (OUT.stem + ".pdf")
    for n, page in enumerate(fitz.open(pdf), 1):
        page.get_pixmap(dpi=150).save(out / f"slide-{n}.png")
    pdf.unlink()
    print(f"Rendered PNGs to {out.relative_to(ROOT)}")


if __name__ == "__main__":
    build(load())
    if "--render" in sys.argv:
        render()
