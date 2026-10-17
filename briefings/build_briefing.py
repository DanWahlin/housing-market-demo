"""Build briefings/housing-market-briefing.pptx from data/ in the GitHub dark template style.

Run: .venv/bin/python briefings/build_briefing.py
"""
import math
from pathlib import Path

import pandas as pd
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.oxml.shapes.autoshape import CT_Shape
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
OUT = ROOT / "briefings" / "housing-market-briefing.pptx"

# Template palette (templates/github-dark-template.pptx)
BG, CARD, BORDER = "0D1117", "161B22", "30363D"
WHITE, BODY, MUTED = "F0F6FC", "B1BAC4", "8B949E"
GREEN, GREEN_DARK, BLUE = "2EA043", "1A5D2A", "58A6FF"
# Color-blind-safe chart colors: blue and orange, plus gray for a third series
C_BLUE, C_ORANGE, C_GRAY = "58A6FF", "F0883E", "8B949E"
FONT = "Arial"
FOOTER = "Built from public FHFA and CFPB data. Not an official Fannie Mae, FHFA, or CFPB product."


def rgb(h):
    return RGBColor.from_string(h)


# ---------- Data ----------
def load():
    m = pd.read_csv(DATA / "hpi_us_division_monthly.csv")
    us = m[m.place_id == "USA"].copy()
    us["quarter"] = (us.month - 1) // 3 + 1
    uq = us.groupby(["year", "quarter"]).agg(v=("index_nsa", "mean"), n=("index_nsa", "size")).reset_index()
    uq = uq[uq.n == 3].reset_index(drop=True)
    uq["yoy"] = uq.v.pct_change(4) * 100

    sq = pd.read_csv(DATA / "hpi_state_quarterly.csv")
    ly, lq = (int(x) for x in sq.sort_values(["year", "quarter"]).iloc[-1][["year", "quarter"]])
    uq = uq[(uq.year < ly) | ((uq.year == ly) & (uq.quarter <= lq))].reset_index(drop=True)

    def uv(y, q):
        return uq[(uq.year == y) & (uq.quarter == q)].v.iloc[0]

    def change(df, key):
        cur = df[(df.year == ly) & (df.quarter == lq)].set_index(key).index_nsa
        prev = df[(df.year == ly - 1) & (df.quarter == lq)].set_index(key).index_nsa
        return (cur / prev * 100 - 100).dropna().sort_values()

    h = pd.read_csv(DATA / "hmda_state_originations.csv")
    peak = uq.loc[uq.yoy.idxmax()]
    return dict(
        ly=ly, lq=lq, period=f"{ly} Q{lq}", uq=uq,
        us_level=uv(ly, lq), us_yoy=uv(ly, lq) / uv(ly - 1, lq) * 100 - 100,
        us_19=uv(ly, lq) / uv(2019, 4) * 100 - 100,
        record=uv(ly, lq) >= uq.v.max(),
        peak_yoy=peak.yoy, peak_label=f"{int(peak.year)} Q{int(peak.quarter)}",
        st=change(sq, "state_code"),
        st_names=pd.read_csv(DATA / "states.csv").set_index("state_code").state_name,
        mt=change(pd.read_csv(DATA / "hpi_metro_quarterly.csv"), "metro_name"),
        hm=h.pivot_table(index="year", columns="loan_purpose", values="loans", aggfunc="sum"),
    )


def short_metro(name):
    city, _, state = name.replace(" (MSAD)", "").partition(", ")
    return f"{city.split('-')[0]}, {state.split('-')[0]}"


# ---------- Drawing helpers ----------
def text(slide, x, y, w, h, value, size=12, color=BODY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    f = r.font
    f.name, f.size, f.bold, f.color.rgb = FONT, Pt(size), bold, rgb(color)
    return tb


def box(slide, x, y, w, h, fill=CARD, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.05):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = adj
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def base_slide(prs, ctx, kicker, headline=None, sub=None, source=None):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    # Top-right pill and GitHub logo, as in the template
    box(s, 7.66, 0.16, 1.68, 0.31, fill=GREEN_DARK, line=GREEN, adj=0.19)
    text(s, 7.66, 0.16, 1.68, 0.31, f"Data through {ctx['period']}", size=10, color="FFFFFF", bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    s.shapes.add_picture(str(ctx["logo"]), Inches(9.48), Inches(0.15), Inches(0.33), Inches(0.33))
    text(s, 0.53, 0.42, 7.0, 0.2, kicker.upper(), size=11, color=GREEN, bold=True)
    if headline:
        text(s, 0.53, 0.66, 8.9, 0.85, headline, size=24, color=WHITE, bold=True)
    if sub:
        text(s, 0.53, 1.6, 8.9, 0.25, sub, size=13, color=BODY)
    if source:
        text(s, 0.53, 4.98, 8.9, 0.2, source, size=9, color=MUTED)
    text(s, 0.53, 5.24, 8.9, 0.2, FOOTER, size=8, color=MUTED)
    return s


def notes(slide, body):
    ns = slide.notes_slide
    if ns.notes_placeholder is None:
        # The template notes master has no body placeholder, so add one to the notes slide
        sp = CT_Shape.new_placeholder_sp(ns.shapes._next_shape_id, "Notes Placeholder", PP_PLACEHOLDER.BODY, "horz", "full", 1)
        ns.shapes._spTree.append(sp)
        ph = ns.notes_placeholder
        ph.left, ph.top, ph.width, ph.height = Inches(0.75), Inches(5.0), Inches(6.0), Inches(4.5)
    ns.notes_text_frame.text = body


def no_fill(parent, before=None):
    spPr = parent.find(qn("c:spPr"))
    if spPr is None:
        spPr = etree.Element(qn("c:spPr"))
        before.addprevious(spPr) if before is not None else parent.append(spPr)
    spPr.clear()
    etree.SubElement(spPr, qn("a:noFill"))
    etree.SubElement(etree.SubElement(spPr, qn("a:ln")), qn("a:noFill"))


def style_chart(chart, legend=False):
    cs = chart._chartSpace
    no_fill(cs, before=cs.find(qn("c:txPr")))
    no_fill(cs.chart.plotArea)
    chart.has_title = False
    # python-pptx writes negative axis IDs; the schema requires unsigned integers
    for el in cs.iter(qn("c:axId"), qn("c:crossAx")):
        el.set("val", str(abs(int(el.get("val")))))
    chart.font.name = FONT
    chart.font.size = Pt(11)
    chart.font.color.rgb = rgb(BODY)
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(11)
        chart.legend.font.color.rgb = rgb(BODY)
    for ax in (chart.category_axis, chart.value_axis):
        ax.format.line.color.rgb = rgb(BORDER)
        ax.tick_labels.font.size = Pt(10)
        ax.tick_labels.font.color.rgb = rgb(BODY)
    chart.value_axis.has_major_gridlines = True
    gl = chart.value_axis.major_gridlines.format.line
    gl.color.rgb = rgb(BORDER)
    gl.width = Pt(0.75)


def label_skip(axis, n):
    """Show every nth category label (schema order: before c:noMultiLvlLbl)."""
    anchor = axis._element.find(qn("c:noMultiLvlLbl"))
    for tag in ("c:tickLblSkip", "c:tickMarkSkip"):
        el = etree.Element(qn(tag), val=str(n))
        anchor.addprevious(el) if anchor is not None else axis._element.append(el)


def axis_title(axis, label):
    axis.has_title = True
    tf = axis.axis_title.text_frame
    tf.text = label
    f = tf.paragraphs[0].runs[0].font
    f.size, f.bold, f.name, f.color.rgb = Pt(10), False, FONT, rgb(MUTED)


def color_point(series, i, color):
    pt = series.points[i]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = rgb(color)
    # LibreOffice draws negative bars as positive unless each data point repeats invertIfNegative
    dPt = series._element.get_or_add_dPt_for_point(i)
    if dPt.find(qn("c:invertIfNegative")) is None:
        dPt.find(qn("c:idx")).addnext(etree.Element(qn("c:invertIfNegative"), val="0"))
    return pt


def point_label(series, i, label, pos, size=11, color=WHITE):
    pt = series.points[i]
    pt.data_label.has_text_frame = True
    pt.data_label.text_frame.text = label
    pt.data_label.position = pos
    for para in pt.data_label.text_frame.paragraphs:
        f = para.runs[0].font
        f.size, f.bold, f.color.rgb = Pt(size), True, rgb(color)
    return pt


def fill_series(series, color):
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = rgb(color)
    series.format.line.fill.background()


def chart_card(slide, title, chart_type, cd, legend=False):
    box(slide, 0.53, 2.0, 8.94, 2.9)
    text(slide, 0.75, 2.1, 8.5, 0.25, title, size=12, color=WHITE, bold=True)
    ch = slide.shapes.add_chart(chart_type, Inches(0.65), Inches(2.35), Inches(8.7), Inches(2.5), cd).chart
    style_chart(ch, legend)
    return ch


def pct_axis(axis, fmt='0"%"'):
    axis.tick_labels.number_format = fmt
    axis.tick_labels.number_format_is_linked = False


# ---------- Slides ----------
def build():
    d = load()
    prs = Presentation(ROOT / "templates" / "github-dark-template.pptx")
    logo = ROOT / "briefings" / ".logo.png"
    logo.write_bytes(next(sh for sh in prs.slides[0].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE).image.blob)
    # Remove the template's sample slides; keep its master, layout, and theme
    for sid in list(prs.slides._sldIdLst):
        prs.part.drop_rel(sid.rId)
        prs.slides._sldIdLst.remove(sid)
    ctx = {"period": d["period"], "logo": logo}

    P, ly, lq, uq = d["period"], d["ly"], d["lq"], d["uq"]
    st, mt, hm, sn = d["st"], d["mt"], d["hm"], d["st_names"]
    hpi_src = f"Source: FHFA House Price Index, purchase-only, not seasonally adjusted (1991 Q1 = 100). Latest period: {P}."
    hmda_src = "Source: CFPB HMDA Data Browser, originated loans, 50 states and DC. Latest year: 2025."
    lead_st, lag_st = st.index[-1], st.index[0]
    st_down, mt_down, mt_n = int((st < 0).sum()), int((mt < 0).sum()), len(mt)

    # 1. Title and KPI cards
    s = base_slide(prs, ctx, "U.S. housing market briefing")
    box(s, 0.53, 1.0, 0.78, 0.07, fill=GREEN, line=None, shape=MSO_SHAPE.RECTANGLE)
    text(s, 0.53, 1.25, 8.9, 1.0,
         f"U.S. home prices are {d['us_19']:.0f}% above 2019 Q4, but annual growth has slowed to {d['us_yoy']:.1f}%",
         size=30, color=WHITE, bold=True)
    text(s, 0.53, 2.35, 8.9, 0.3,
         "Where prices are, how fast they change, who leads, and how mortgage lending changed.", size=14)
    kpis = [
        (f"{d['us_level']:.1f}", "U.S. house price index", f"{P}, 1991 Q1 = 100"),
        (f"+{d['us_yoy']:.1f}%", "U.S. year-over-year change", f"{P} compared with {ly - 1} Q{lq}"),
        (f"+{d['us_19']:.1f}%", "Change since 2019 Q4", "Last quarter before the pandemic"),
    ]
    for i, (big, lab, sub) in enumerate(kpis):
        x = 0.53 + i * 3.07
        box(s, x, 2.95, 2.8, 1.75)
        box(s, x, 2.95, 2.8, 0.06, fill=BLUE, line=None, shape=MSO_SHAPE.RECTANGLE)
        text(s, x + 0.22, 3.2, 2.4, 0.6, big, size=32, color=WHITE, bold=True)
        text(s, x + 0.22, 3.85, 2.4, 0.3, lab, size=12, color=BLUE, bold=True)
        text(s, x + 0.22, 4.15, 2.4, 0.4, sub, size=10, color=MUTED)
    text(s, 0.53, 4.98, 8.9, 0.2, hpi_src, size=9, color=MUTED)
    notes(s, f"This briefing uses public FHFA house price index data through {P} and CFPB HMDA mortgage data for 2018 to 2025. "
             f"The U.S. index is {d['us_level']:.1f}, where 1991 Q1 equals 100. The index shows price change, not dollar prices. "
             f"Prices are {d['us_19']:.1f}% above 2019 Q4, the last quarter before the pandemic. "
             f"Compared with {ly - 1} Q{lq}, they rose {d['us_yoy']:.1f}%. "
             "Each U.S. quarterly value is the average of the three monthly values.")

    # 2. Where prices are now: long-run index
    rec = "a record high" if d["record"] else "near its high"
    s = base_slide(prs, ctx, "Where prices are now",
                   f"The U.S. price index is at {rec}, {d['us_19']:.0f}% above its pre-pandemic level",
                   f"Quarterly U.S. house price index, 1991 Q1 to {P}. Most of the gain since 2019 came from 2020 to 2022.",
                   hpi_src)
    cd = CategoryChartData()
    cd.categories = [str(y) if q == 1 else f"{y} Q{q}" for y, q in zip(uq.year, uq.quarter)]
    cd.add_series("U.S. HPI", [round(v, 1) for v in uq.v])
    ch = chart_card(s, "U.S. house price index (1991 Q1 = 100), quarterly", XL_CHART_TYPE.LINE, cd)
    ser = ch.series[0]
    ser.format.line.color.rgb = rgb(C_BLUE)
    ser.format.line.width = Pt(2.5)
    ser.smooth = False
    label_skip(ch.category_axis, 20)
    ch.value_axis.minimum_scale, ch.value_axis.maximum_scale = 0, 600
    pct_axis(ch.value_axis, "0")
    axis_title(ch.value_axis, "Index level")
    axis_title(ch.category_axis, "Year")
    point_label(ser, len(uq) - 1, f"{P}\n{d['us_level']:.1f}", XL_LABEL_POSITION.ABOVE)
    i19 = int(uq.index[(uq.year == 2019) & (uq.quarter == 4)][0])
    point_label(ser, i19, f"2019 Q4\n{uq.v[i19]:.1f}", XL_LABEL_POSITION.ABOVE, size=10, color=BODY)
    notes(s, "The line shows the U.S. purchase-only house price index since 1991. "
             f"In 2019 Q4 the index was {uq.v[i19]:.1f}. In {P} it is {d['us_level']:.1f}, an increase of {d['us_19']:.1f}%. "
             "The steepest climb was from 2020 to 2022. Since then the line continues to go up, but more slowly. "
             "This is an index of price change, not an average dollar price.")

    # 3. How fast prices changed: YoY by quarter since 2018
    rq = uq[uq.year >= 2018].reset_index(drop=True)
    s = base_slide(prs, ctx, "How fast prices changed",
                   f"Annual price growth fell from {d['peak_yoy']:.1f}% in {d['peak_label']} to {d['us_yoy']:.1f}% in {P}",
                   "Year-over-year change in the U.S. index, by quarter. Prices still rise, but at a much lower rate.",
                   hpi_src)
    cd = CategoryChartData()
    cd.categories = [f"{y} Q{q}" for y, q in zip(rq.year, rq.quarter)]
    cd.add_series("Year-over-year change (%)", [round(v, 1) for v in rq.yoy])
    ch = chart_card(s, "U.S. house prices: change from the same quarter one year earlier (%)", XL_CHART_TYPE.COLUMN_CLUSTERED, cd)
    ser = ch.series[0]
    fill_series(ser, C_BLUE)
    ch.plots[0].gap_width = 40
    for i, lab in enumerate(cd.categories):
        if lab.label in (d["peak_label"], P):
            color_point(ser, i, C_ORANGE)
            point_label(ser, i, f"{rq.yoy[i]:.1f}%", XL_LABEL_POSITION.OUTSIDE_END)
    label_skip(ch.category_axis, 4)
    pct_axis(ch.value_axis)
    ch.value_axis.minimum_scale, ch.value_axis.maximum_scale = 0, 22
    axis_title(ch.value_axis, "Change (%)")
    axis_title(ch.category_axis, "Quarter")
    notes(s, "Each bar shows how much U.S. prices changed compared with the same quarter one year earlier. "
             f"Growth peaked at {d['peak_yoy']:.1f}% in {d['peak_label']}. In {P} it is {d['us_yoy']:.1f}%. "
             "The two orange bars mark the peak and the latest quarter. "
             "At the national level prices did not fall; the rate of increase became smaller.")

    # 4 and 5. Leaders and laggards
    def rank_slide(kicker, headline, sub, series, label, title, note):
        s = base_slide(prs, ctx, kicker, headline, sub, hpi_src)
        sel = pd.concat([series.head(5), series.tail(5)])
        cd = CategoryChartData()
        cd.categories = [label(k) for k in sel.index]
        cd.add_series("Year-over-year change (%)", [round(v, 1) for v in sel.values])
        ch = chart_card(s, title, XL_CHART_TYPE.BAR_CLUSTERED, cd)
        ser = ch.series[0]
        fill_series(ser, C_BLUE)
        ser.invert_if_negative = False
        ch.plots[0].gap_width = 35
        for i, v in enumerate(sel.values):
            if v < 0:
                color_point(ser, i, C_ORANGE)
        plot = ch.plots[0]
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format, dl.number_format_is_linked = '+0.0"%";-0.0"%"', False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(9)
        dl.font.color.rgb = rgb(WHITE)
        ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        pct_axis(ch.value_axis)
        label_skip(ch.category_axis, 1)
        ch.category_axis.tick_labels.font.size = Pt(9)
        # Extra room on the left so negative data labels do not touch the category names
        ch.value_axis.minimum_scale = math.floor((min(sel.min(), 0) - 1) / 2) * 2
        ch.value_axis.maximum_scale = math.ceil(sel.max() / 2) * 2 + 2
        ch.value_axis.major_unit = 2
        axis_title(ch.value_axis, f"Change from {ly - 1} Q{lq} to {P} (%)")
        notes(s, note)

    fmt = lambda items, name: ", ".join(f"{name(k)} {v:+.1f}%" for k, v in items)
    rank_slide(
        "Leaders and laggards: states",
        f"{sn[lead_st]} leads at +{st.iloc[-1]:.1f}%; {st_down} states, with {sn[lag_st]} last, show price declines",
        f"Five fastest and five slowest of 50 states and DC. U.S.: +{d['us_yoy']:.1f}%. Blue = increase, orange = decrease.",
        st, lambda k: sn[k], f"State house price index: year-over-year change, {P}",
        f"These are the five states with the largest and the five with the smallest year-over-year change in {P}. "
        f"Top five: {fmt(st.tail(5)[::-1].items(), lambda k: sn[k])}. "
        f"Bottom five: {fmt(st.head(5).items(), lambda k: sn[k])}. "
        f"{st_down} of {len(st)} states and DC had a decrease. Four of the five slowest states are in the West.")
    rank_slide(
        "Leaders and laggards: metros",
        f"{mt_down} of the {mt_n} largest metros show price declines; {short_metro(mt.index[-1])} leads at +{mt.iloc[-1]:.1f}%",
        f"Five fastest and five slowest of the {mt_n} largest metros. Blue = increase, orange = decrease.",
        mt, short_metro, f"Metro house price index: year-over-year change, {P}",
        f"These are the five metros with the largest and the five with the smallest year-over-year change in {P}, "
        f"out of the {mt_n} largest metros in the FHFA file. Some are metro divisions. Labels use the first city name. "
        f"Top five: {fmt(mt.tail(5)[::-1].items(), str)}. "
        f"Bottom five: {fmt(mt.head(5).items(), str)}. "
        f"{mt_down} of {mt_n} metros had a decrease.")

    # 6. Originations 2018 to 2025
    pur, refi, co = hm["Home purchase"], hm["Refinance"], hm["Cash-out refinance"]
    allrefi = refi + co
    pk = int(allrefi.idxmax())
    refi_drop = 100 - allrefi[2025] / allrefi[pk] * 100
    pur_drop = 100 - pur[2025] / pur[2018] * 100
    s = base_slide(prs, ctx, "Mortgage originations, 2018 to 2025",
                   f"Refinance loans fell {refi_drop:.0f}% from the {pk} peak; home purchase loans are {pur_drop:.0f}% below 2018",
                   "Originated loans each year by loan purpose. Refinance figures in the headline include cash-out refinance.",
                   hmda_src)
    cd = CategoryChartData()
    cd.categories = [str(y) for y in hm.index]
    for name, vals in (("Home purchase", pur), ("Refinance", refi), ("Cash-out refinance", co)):
        cd.add_series(name, [round(v / 1e6, 2) for v in vals])
    ch = chart_card(s, "Originated mortgage loans by purpose (millions)", XL_CHART_TYPE.COLUMN_CLUSTERED, cd, legend=True)
    for ser_, c in zip(ch.series, (C_BLUE, C_ORANGE, C_GRAY)):
        fill_series(ser_, c)
    ch.plots[0].gap_width, ch.plots[0].overlap = 60, 0
    pct_axis(ch.value_axis, "0.0")
    ch.value_axis.minimum_scale = 0
    axis_title(ch.value_axis, "Loans (millions)")
    axis_title(ch.category_axis, "HMDA reporting year")
    notes(s, "This chart uses CFPB HMDA data for originated loans only. "
             f"Home purchase loans: {pur[2018] / 1e6:.2f} million in 2018, a high of {pur.max() / 1e6:.2f} million in {int(pur.idxmax())}, "
             f"and {pur[2025] / 1e6:.2f} million in 2025. "
             f"Refinance plus cash-out refinance: {allrefi[2018] / 1e6:.2f} million in 2018, {allrefi[pk] / 1e6:.2f} million at the {pk} peak, "
             f"{allrefi.min() / 1e6:.2f} million at the low in {int(allrefi.idxmin())}, and {allrefi[2025] / 1e6:.2f} million in 2025. "
             "In 2020 and 2021 there were more refinance loans than home purchase loans. Since 2022 home purchase loans are the larger group.")

    # 7. Three takeaways
    s = base_slide(prs, ctx, "Three takeaways",
                   "Prices are high and still rise slowly, but gains are uneven and refinance activity is low",
                   "Facts from the FHFA and CFPB data in this briefing. No forecasts.",
                   f"Sources: FHFA HPI (latest period {P}); CFPB HMDA (2018 to 2025).")
    cards = [
        ("Prices are high", f"+{d['us_19']:.0f}%",
         f"U.S. index change since 2019 Q4. Growth slowed to +{d['us_yoy']:.1f}% year over year in {P}.", C_BLUE),
        ("Gains are uneven", f"{mt_down} of {mt_n}",
         f"largest metros and {st_down} states show price declines. {sn[lead_st]} leads states at +{st.iloc[-1]:.1f}%.", C_ORANGE),
        ("Refinance is low", f"-{refi_drop:.0f}%",
         f"Refinance loans in 2025 compared with {pk}. Home purchase loans are {pur_drop:.0f}% below 2018.", C_GRAY),
    ]
    for i, (head, big, body, c) in enumerate(cards):
        x = 0.53 + i * 3.07
        box(s, x, 1.95, 2.8, 2.95)
        box(s, x, 1.95, 2.8, 0.06, fill=c, line=None, shape=MSO_SHAPE.RECTANGLE)
        box(s, x + 0.22, 2.2, 0.32, 0.32, fill=BLUE, line=None, shape=MSO_SHAPE.OVAL)
        text(s, x + 0.22, 2.2, 0.32, 0.32, str(i + 1), size=12, color="FFFFFF", bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.65, 2.2, 2.0, 0.32, head, size=13, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.22, 2.75, 2.4, 0.6, big, size=28, color=WHITE, bold=True)
        text(s, x + 0.22, 3.45, 2.4, 1.35, body, size=12)
    notes(s, f"First, U.S. prices are {d['us_19']:.1f}% above 2019 Q4 and rose {d['us_yoy']:.1f}% in the last year. "
             f"Second, the change is uneven: {st_down} states and {mt_down} of the {mt_n} largest metros show a decrease, "
             f"while {sn[lead_st]} rose {st.iloc[-1]:.1f}%. "
             f"Third, refinance and cash-out refinance originations in 2025 were {refi_drop:.0f}% below the {pk} peak, "
             f"and home purchase originations were {pur_drop:.0f}% below 2018. "
             "These are facts from the data, not forecasts.")

    prs.save(OUT)
    logo.unlink()
    print(f"Saved {OUT}")


if __name__ == "__main__":
    build()
